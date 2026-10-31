"""Generate the blank import workbook. Requires openpyxl."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
import re
from pathlib import Path

fields = [
('fullName','Nama lengkap (wajib)'),('nickname','Nama panggilan'),('nik','NIK 16 digit, teks'),('familyCardNo','Nomor KK kependudukan, teks'),('familyNo','Nomor keluarga/KK aplikasi yang sudah ada (bukan nomor KK kependudukan)'),('familyRole','Hubungan dalam keluarga'),('gender','Jenis kelamin (wajib)'),('birthPlace','Tempat lahir'),('birthDate','Tanggal lahir (wajib)'),('address','Alamat'),('kelurahan','Kelurahan/desa'),('kecamatan','Kecamatan'),('kabupaten','Kabupaten/kota'),('phone','Nomor HP, teks'),('email','Email'),('maritalStatus','Status perkawinan'),('spouseName','Nama pasangan'),('rayon','Nama, kode atau ID rayon yang sudah ada (wajib)'),('categorial','Nama, kode atau ID kategorial yang sudah ada'),('membershipStatus','Status keanggotaan'),('baptismStatus','Status baptis'),('baptismDate','Tanggal baptis'),('confirmationStatus','Status sidi'),('confirmationDate','Tanggal sidi'),('marriageChurchStatus','Status nikah gereja'),('marriageDate','Tanggal nikah gereja'),('education','Pendidikan'),('occupation','Pekerjaan'),('joinDate','Tanggal bergabung'),('status','Status jemaat'),('notes','Catatan')]
choices = {'familyRole':['Kepala Keluarga','Istri','Anak','Orang Tua','Famili Lain'], 'gender':['Laki-laki','Perempuan'], 'maritalStatus':['Belum Menikah','Menikah','Duda','Janda'], 'membershipStatus':['Tetap','Tamu','Titipan'], 'baptismStatus':['Sudah','Belum'], 'confirmationStatus':['Sudah','Belum'], 'marriageChurchStatus':['Sudah','Belum'], 'status':['Aktif','Pindah','Meninggal']}
types = Path('src/types/index.ts').read_text()
for key, constant in [('education','STANDARD_EDUCATION_OPTIONS'),('occupation','STANDARD_OCCUPATION_OPTIONS')]:
    block = types.split('export const '+constant+' = [',1)[1].split('] as const',1)[0]
    choices[key] = re.findall(r"'([^']+)'", block)
w=Workbook(); data=w.active; data.title='Data Jemaat'; guide=w.create_sheet('Petunjuk'); lists=w.create_sheet('Pilihan')
data.append([k for k,_ in fields]); data.freeze_panes='C2'; data.auto_filter.ref='A1:AE501'; data.row_dimensions[1].height=34
for i,(key,desc) in enumerate(fields,1):
    c=data.cell(1,i); c.fill=PatternFill('solid',fgColor='174665'); c.font=Font(color='FFFFFF',bold=True); c.alignment=Alignment(wrap_text=True,vertical='center'); c.comment=Comment(desc,'GMIT'); data.column_dimensions[get_column_letter(i)].width=22
    for row in range(2,502): data.cell(row,i).number_format='@'
for col,(key,values) in enumerate(choices.items(),1):
    lists.cell(1,col,key)
    for row,value in enumerate(values,2): lists.cell(row,col,value)
    letter=get_column_letter(col); lists.column_dimensions[letter].width=38
    dv=DataValidation(type='list',formula1=f"'Pilihan'!${letter}$2:${letter}${len(values)+1}",allow_blank=True); dv.errorTitle='Pilihan tidak valid'; dv.error='Pilih dari daftar.'; dv.showErrorMessage=True; data.add_data_validation(dv)
    idx=[k for k,_ in fields].index(key)+1; dv.add(f'{get_column_letter(idx)}2:{get_column_letter(idx)}501')
lines=['TEMPLATE IMPOR DATA JEMAAT', '1. Isi satu orang per baris pada sheet Data Jemaat; jangan mengubah nama sheet atau judul kolom.', '2. Wajib: fullName (minimal 2 huruf), gender, birthDate dan rayon. Maksimal 500 orang per file.', '3. NIK, nomor KK dan nomor HP adalah TEKS. Pertahankan angka nol di depan; NIK harus tepat 16 digit. Jangan memasukkan sebagai angka karena Excel dapat mengubah digit.', '4. Semua tanggal: YYYY-MM-DD, contoh pola 2026-01-31. Tanggal Excel juga diterima.', '5. rayon dan categorial: isi nama, kode atau ID yang sudah ada di aplikasi. Jika nama tidak unik, gunakan kode/ID.', '6. familyNo adalah nomor keluarga yang dibuat aplikasi (mis. KK-BF-...), bukan nomor KK kependudukan. Buat keluarga dahulu; kosongkan jika belum ada.', '7. Pilih nilai dropdown untuk kolom pilihan. Jika kosong: hubungan Anak, status perkawinan Belum Menikah, keanggotaan Tetap, baptis/sidi/nikah Belum, status Aktif; tanggal bergabung otomatis tanggal impor.', '8. Simpan .xlsx, buka menu Impor Data Jemaat di aplikasi dan pilih file. Periksa pratinjau lalu impor. Baris yang gagal ditampilkan; baris berhasil sudah tersimpan. Jangan unggah ulang baris berhasil.', '9. Nomor jemaat dibuat otomatis. Foto dan dokumen tidak diimpor dari Excel; unggah melalui aplikasi secara terpisah.', '10. Template kosong tanpa data contoh. Simpan data pribadi secara aman.', '', 'KOLOM | KETERANGAN']
for x in lines: guide.append([x])
for key,desc in fields: guide.append([key,desc])
guide.column_dimensions['A'].width=35; guide.column_dimensions['B'].width=85
for i in range(1,len(lines)+1): guide.merge_cells(start_row=i,start_column=1,end_row=i,end_column=2); guide.cell(i,1).alignment=Alignment(wrap_text=True,vertical='center'); guide.row_dimensions[i].height=32
for sheet in [guide,lists]:
    for cell in sheet[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='174665')
w.save('public/Template_Impor_Data_Jemaat.xlsx')
